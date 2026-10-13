"""Document text extraction helpers."""

import csv
import datetime
import xml.etree.ElementTree as ElementTree
from dataclasses import dataclass
from html.parser import HTMLParser
from pathlib import Path

import openpyxl
import xlrd
from openpyxl.utils import get_column_letter

from src.utils.reading.ocr import ocr_pdf

# How deep any one sheet is read. A sheet with more rows than this comes
# back ending in row_cap_marker() instead of its last rows.
SHEET_ROW_CAP = 500


def _decode(raw: bytes) -> str:
    """Bytes as text, trying the encodings office software actually writes."""

    for encoding in ("utf-8-sig", "utf-16"):
        try:
            return raw.decode(encoding)
        except UnicodeDecodeError:
            continue
    return raw.decode("latin-1", "replace")


# Left to choose freely, the sniffer picks whatever character is most regular
# across the sample. On a Windows file opened with newline="" — which the csv
# module requires — that can be the "\r" of every line ending, and csv.reader
# then refuses it ("bad delimiter value"). On a file whose first line is a
# company name padded with commas it picked the SPACE, and split
# "CÔNG TY TNHH AN CƯỜNG" into five cells.
CSV_DELIMITERS = ",;\t|"


def _read_csv_rows(csv_path: str) -> list[list[str]]:
    """Every row of a CSV, decoded and stripped. Never truncates."""

    # A byte-order mark is the only thing that says UTF-16 for certain, and
    # guessing it instead would decode an ordinary Latin-1 file into mojibake.
    # Excel's own "Unicode Text" export writes one.
    with open(csv_path, "rb") as probe:
        wide = probe.read(2) in (b"\xff\xfe", b"\xfe\xff")

    last_error = None
    for encoding in (("utf-16", "utf-8", "latin-1") if wide
                     else ("utf-8", "latin-1")):
        try:
            with open(csv_path, newline="", encoding=encoding) as csv_file:
                sample = csv_file.read(4096)
                csv_file.seek(0)
                try:
                    dialect = csv.Sniffer().sniff(sample, delimiters=CSV_DELIMITERS)
                except csv.Error:
                    dialect = csv.excel
                return [
                    [cell.strip() for cell in row]
                    for row in csv.reader(csv_file, dialect)
                ]
        except UnicodeDecodeError as exc:
            last_error = exc
    raise ValueError(f"Unable to decode CSV file: {last_error}")


def extract_csv_text(
    csv_path: str,
    max_rows: int = SHEET_ROW_CAP,
    max_chars: int = 100000,
) -> str:
    """Extract CSV content as tab-separated text for LLM analysis."""

    rows, duplicates = _dedup_rows(_read_csv_rows(csv_path))
    body = rows[:max_rows]
    if len(rows) > max_rows:
        body = body + [[row_cap_marker(max_rows)]]
    if duplicates:
        body = [[f"... đã bỏ {duplicates} dòng trùng lặp ..."]] + body
    return "\n".join("\t".join(row) for row in body)[:max_chars]


def _format_number(value: float, number_format: str) -> str:
    """Render a numeric cell the way Excel displays it."""
    fmt = number_format or ""

    if "%" in fmt:
        scaled = value * 100
        text = f"{scaled:,.2f}".rstrip("0").rstrip(".")
        return _to_vietnamese_number(text or "0") + "%"

    # "#,##0" / "#,##0.00" ask Excel to group thousands.
    if "#,##" in fmt or "0,000" in fmt:
        decimals = 0
        if "." in fmt:
            decimals = len(fmt.split(".")[-1].split("_")[0].rstrip("%);\\ "))
        return _to_vietnamese_number(f"{value:,.{decimals}f}")

    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value)


def _to_vietnamese_number(us_formatted: str) -> str:
    """Swap US separators for Vietnamese ones ('.' thousands, ',' decimal)."""
    return us_formatted.replace(",", "\x00").replace(".", ",").replace("\x00", ".")


def _format_cell(value: object, number_format: str) -> str:
    """Render one cell value as faithful, LLM-readable text."""
    if value is None:
        return ""

    if isinstance(value, datetime.datetime):
        if (value.hour, value.minute, value.second) == (0, 0, 0):
            return value.strftime("%Y-%m-%d")
        return value.strftime("%Y-%m-%d %H:%M")
    if isinstance(value, datetime.date):
        return value.strftime("%Y-%m-%d")
    if isinstance(value, datetime.time):
        return value.strftime("%H:%M")

    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, (int, float)):
        return _format_number(value, number_format)

    return str(value).strip()


def _propagate_merges(
    grid: list[list[str]],
    merges: list[tuple[int, int, int, int]],
) -> list[list[str]]:
    """Copy each merged block's value into every cell the block covers.

    Boxes are (top, left, bottom, right), zero-based, bottom and right
    exclusive. A box that starts in the first column and spans the whole used
    width is left where it is: that shape is a banner — a period line, a company
    name — and spreading it sideways makes the row read like data.

    Every engine below hands its merges here so the four cannot drift apart:
    whether a header is one cell or four is what tells a reader which column a
    figure belongs to.
    """

    if not grid:
        return []

    used_width = max((len(row) for row in grid), default=0)
    for top, left, bottom, right in merges:
        if top >= len(grid) or left >= len(grid[top]):
            continue
        value = grid[top][left]
        if not value:
            continue
        if left == 0 and right >= used_width:
            continue
        for row_index in range(top, min(bottom, len(grid))):
            for col_index in range(left, right):
                if col_index < len(grid[row_index]):
                    grid[row_index][col_index] = value
    return grid


def _sheet_grid(worksheet) -> list[list[str]]:
    """Return the sheet as rendered text cells, with merged values propagated."""
    grid = [
        [_format_cell(cell.value, cell.number_format) for cell in row]
        for row in worksheet.iter_rows()
    ]
    return _propagate_merges(
        grid,
        [
            (merged.min_row - 1, merged.min_col - 1, merged.max_row, merged.max_col)
            for merged in worksheet.merged_cells.ranges
        ],
    )


def _dense_grid(cells: dict[tuple[int, int], str]) -> list[list[str]]:
    """A sparse {(row, column): text} map as a rectangular grid."""

    if not cells:
        return []
    height = max(row for row, _ in cells) + 1
    width = max(column for _, column in cells) + 1
    return [
        [cells.get((row, column), "") for column in range(width)]
        for row in range(height)
    ]


def _clean_grid(grid: list[list[str]]) -> tuple[list[list[str]], list[int]]:
    """Drop all-empty rows/columns; return the grid and its kept column indexes."""
    if not grid:
        return [], []

    width = max(len(row) for row in grid)
    padded = [row + [""] * (width - len(row)) for row in grid]
    kept_columns = [
        index for index in range(width) if any(row[index] for row in padded)
    ]
    cleaned = [
        [row[index] for index in kept_columns]
        for row in padded
        if any(row[index] for index in kept_columns)
    ]
    return cleaned, kept_columns


def _dedup_rows(grid: list[list[str]]) -> tuple[list[list[str]], int]:
    """Keep the first occurrence of each distinct row; also return how many went."""

    seen: set[tuple[str, ...]] = set()
    kept: list[list[str]] = []
    for row in grid:
        key = tuple(row)
        if key in seen:
            continue
        seen.add(key)
        kept.append(row)
    return kept, len(grid) - len(kept)


def _sheet_header(title: str, rows: int, columns: int, duplicates: int) -> str:
    """The "--- Sheet: … ---" line, naming duplicates when any were dropped."""

    dropped = f", đã bỏ {duplicates} dòng trùng lặp" if duplicates else ""
    return f"--- Sheet: {title} ({rows} dòng x {columns} cột{dropped}) ---"


def row_cap_marker(max_rows: int) -> str:
    """The line that stands in for the rows a sheet was cut off at.

    Named here so the ledger pass can recognise its own reader's cut rather
    than matching the wording by hand: a sheet read down to 500 rows of 4,000
    has no usable total and no usable row count, and saying nothing about that
    is how a quarter of a ledger gets reported as the whole of it.
    """

    return f"... truncated after {max_rows} rows ..."


def _grid_to_tsv(
    grid: list[list[str]],
    kept_columns: list[int],
    max_rows: int,
) -> str:
    """Render a cleaned grid as TSV led by its Excel column letters."""
    lines = ["\t".join(get_column_letter(index + 1) for index in kept_columns)]
    for row in grid[:max_rows]:
        # Right-trim trailing empties so short rows stop emitting runs of tabs.
        trimmed = list(row)
        while trimmed and not trimmed[-1]:
            trimmed.pop()
        lines.append("\t".join(trimmed))
    if len(grid) > max_rows:
        lines.append(row_cap_marker(max_rows))
    return "\n".join(lines)


@dataclass(frozen=True)
class SheetText:
    """One sheet of a spreadsheet, already cleaned, de-duplicated and capped."""

    sheet_name: str
    content: str
    row_count: int
    column_count: int
    duplicates_dropped: int


def _sheet_text(
    title: str,
    grid: list[list[str]],
    max_rows_per_sheet: int,
) -> SheetText:
    """One engine's grid, cleaned, de-duplicated and rendered as TSV."""

    cleaned, kept_columns = _clean_grid(grid)
    if not cleaned:
        return SheetText(title, "", 0, 0, 0)
    cleaned, duplicates = _dedup_rows(cleaned)
    return SheetText(
        title,
        _grid_to_tsv(cleaned, kept_columns, max_rows_per_sheet),
        len(cleaned),
        len(kept_columns),
        duplicates,
    )


_ZIP_SIGNATURE = b"PK\x03\x04"
_OLE2_SIGNATURE = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"
# A raw BIFF record stream opens with the BOF record, 0x0009, and the second
# byte says which BIFF version wrote it.
_BIFF_SIGNATURES = (b"\x09\x00", b"\x09\x02", b"\x09\x04", b"\x09\x08")
_SPREADSHEETML_MARKS = (
    b"urn:schemas-microsoft-com:office:spreadsheet",
    b'progid="excel.sheet"',
)
_HTML_MARKS = (b"<table", b"<html", b"<!doctype html")


def _workbook_shape(path: str) -> str:
    """What the file actually is. The extension is a hint, not an answer.

    ".xls" is the extension four unrelated formats are saved under: the binary
    Excel 97-2003 workbook, an HTML table, Excel 2003 XML, and a plain .xlsx
    somebody renamed. Accounting software here emits the middle two routinely.
    Dispatching on the extension sent all four to one engine, which could open
    only the first, and the other three came back as a failed extraction.
    """

    with open(path, "rb") as handle:
        head = handle.read(8192)

    if head.startswith(_ZIP_SIGNATURE):
        return "xlsx"
    if head.startswith(_OLE2_SIGNATURE) or head[:2] in _BIFF_SIGNATURES:
        return "xls"
    lowered = head.lower()
    if any(mark in lowered for mark in _SPREADSHEETML_MARKS):
        return "spreadsheetml"
    if any(mark in lowered for mark in _HTML_MARKS):
        return "html"
    return "text"


def read_sheets(
    excel_path: str,
    max_rows_per_sheet: int = SHEET_ROW_CAP,
) -> list[SheetText]:
    """Every sheet of a spreadsheet, one record each, whatever wrote the file."""

    shape = _workbook_shape(excel_path)
    if shape == "xlsx":
        return _read_xlsx_sheets(excel_path, max_rows_per_sheet)
    if shape == "xls":
        return _read_xls_sheets(excel_path, max_rows_per_sheet)
    if shape == "spreadsheetml":
        return _read_spreadsheetml_sheets(excel_path, max_rows_per_sheet)
    if shape == "html":
        return _read_html_table_sheets(excel_path, max_rows_per_sheet)

    rows, duplicates = _dedup_rows(_read_csv_rows(excel_path))
    grid, kept_columns = _clean_grid(rows)
    if not grid:
        return [SheetText("", "", 0, 0, duplicates)]
    return [SheetText("", _grid_to_tsv(grid, kept_columns, max_rows_per_sheet),
                      len(grid), len(kept_columns), duplicates)]


def _read_xlsx_sheets(
    excel_path: str,
    max_rows_per_sheet: int,
) -> list[SheetText]:
    """An Open XML workbook, read with openpyxl.

    Opened as a stream rather than by name: given a path, openpyxl decides what
    the file is from its extension and refuses an .xlsx that someone saved as
    .xls, which is a file it can read perfectly well.
    """

    with open(excel_path, "rb") as handle:
        workbook = openpyxl.load_workbook(handle, data_only=True)
        try:
            return [
                _sheet_text(
                    worksheet.title, _sheet_grid(worksheet), max_rows_per_sheet
                )
                for worksheet in workbook.worksheets
            ]
        finally:
            workbook.close()


def _read_xls_sheets(
    excel_path: str,
    max_rows_per_sheet: int,
) -> list[SheetText]:
    """A binary Excel workbook (BIFF), read with xlrd.

    ``formatting_info`` is what carries the number formats and the merged
    ranges, so asking for it is what makes an .xls ledger read like the .xlsx
    path rather than a flattened version of it. Some workbooks make xlrd refuse
    it, and for those the plain read still returns every value.
    """

    try:
        book = xlrd.open_workbook(excel_path, formatting_info=True)
    except Exception:
        book = xlrd.open_workbook(excel_path)
    try:
        return [
            _sheet_text(sheet.name, _xls_grid(book, sheet), max_rows_per_sheet)
            for sheet in book.sheets()
        ]
    finally:
        book.release_resources()


def _xls_number_format(book, cell) -> str:
    """The cell's Excel format string, or "" when the workbook withheld it.

    Without ``formatting_info`` there is no style index to follow, and a few
    workbooks carry indexes that point at nothing. Either way the value is
    still rendered, just without thousands grouping.
    """

    try:
        style = book.xf_list[cell.xf_index]
        return book.format_map[style.format_key].format_str or ""
    except Exception:
        return ""


def _xls_cell(book, cell, datemode: int) -> str:
    """One BIFF cell as text, matching what the .xlsx path would print."""

    if cell.ctype in (xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK):
        return ""
    if cell.ctype == xlrd.XL_CELL_ERROR:
        return xlrd.error_text_from_code.get(cell.value, "")
    if cell.ctype == xlrd.XL_CELL_BOOLEAN:
        return "TRUE" if cell.value else "FALSE"
    if cell.ctype == xlrd.XL_CELL_DATE:
        stamp = xlrd.xldate.xldate_as_datetime(cell.value, datemode)
        return _format_cell(stamp.time() if cell.value < 1 else stamp, "")
    if cell.ctype == xlrd.XL_CELL_NUMBER:
        return _format_number(cell.value, _xls_number_format(book, cell))
    return str(cell.value).strip()


def _xls_grid(book, sheet) -> list[list[str]]:
    """One BIFF sheet as rendered text cells, with merged values propagated."""

    grid = [
        [
            _xls_cell(book, sheet.cell(row, column), book.datemode)
            for column in range(sheet.ncols)
        ]
        for row in range(sheet.nrows)
    ]
    return _propagate_merges(
        grid,
        [
            (top, left, bottom, right)
            for top, bottom, left, right in sheet.merged_cells
        ],
    )


_SSML = "{urn:schemas-microsoft-com:office:spreadsheet}"


def _ssml_number_formats(root) -> dict[str, str]:
    """{StyleID: Excel format string} from the workbook's <Styles> block."""

    formats = {}
    for style in root.iter(f"{_SSML}Style"):
        style_id = style.get(f"{_SSML}ID")
        number_format = style.find(f"{_SSML}NumberFormat")
        if style_id and number_format is not None:
            formats[style_id] = number_format.get(f"{_SSML}Format", "")
    return formats


def _ssml_value(cell, formats: dict[str, str]) -> str:
    """One Excel-2003-XML cell as text."""

    data = cell.find(f"{_SSML}Data")
    if data is None:
        return ""
    text = "".join(data.itertext()).strip()
    kind = data.get(f"{_SSML}Type", "String")
    if kind == "Number":
        try:
            value = float(text)
        except ValueError:
            return text
        return _format_number(
            value, formats.get(cell.get(f"{_SSML}StyleID", ""), "")
        )
    if kind == "DateTime":
        try:
            return _format_cell(datetime.datetime.fromisoformat(text), "")
        except ValueError:
            return text
    if kind == "Boolean":
        return "TRUE" if text not in ("0", "") else "FALSE"
    return text


def _ssml_grid(table, formats: dict[str, str]):
    """An Excel-2003-XML <Table> as (grid, merge boxes).

    ``ss:Index`` is how the format skips empty rows and columns, so a reader
    that just counts the elements it sees shifts every figure after the first
    gap into the wrong column.
    """

    cells: dict[tuple[int, int], str] = {}
    merges: list[tuple[int, int, int, int]] = []
    row_index = -1
    for row in table.findall(f"{_SSML}Row"):
        declared_row = row.get(f"{_SSML}Index")
        row_index = int(declared_row) - 1 if declared_row else row_index + 1
        column_index = 0
        for cell in row.findall(f"{_SSML}Cell"):
            declared = cell.get(f"{_SSML}Index")
            if declared:
                column_index = int(declared) - 1
            across = int(cell.get(f"{_SSML}MergeAcross", 0))
            down = int(cell.get(f"{_SSML}MergeDown", 0))
            cells[(row_index, column_index)] = _ssml_value(cell, formats)
            if across or down:
                merges.append((
                    row_index,
                    column_index,
                    row_index + down + 1,
                    column_index + across + 1,
                ))
            column_index += across + 1
    return _dense_grid(cells), merges


def _read_spreadsheetml_sheets(
    excel_path: str,
    max_rows_per_sheet: int,
) -> list[SheetText]:
    """An Excel 2003 XML workbook, which is normally saved as ".xls"."""

    raw = Path(excel_path).read_text(encoding="utf-8-sig", errors="replace")
    root = ElementTree.fromstring(raw)
    formats = _ssml_number_formats(root)

    sheets = []
    for worksheet in root.iter(f"{_SSML}Worksheet"):
        title = worksheet.get(f"{_SSML}Name", "")
        table = worksheet.find(f"{_SSML}Table")
        if table is None:
            sheets.append(SheetText(title, "", 0, 0, 0))
            continue
        grid, merges = _ssml_grid(table, formats)
        sheets.append(_sheet_text(
            title, _propagate_merges(grid, merges), max_rows_per_sheet,
        ))
    return sheets


class _HtmlTableReader(HTMLParser):
    """Every <table> in an HTML file, as rows of (text, colspan, rowspan)."""

    def __init__(self) -> None:
        super().__init__(convert_charrefs=True)
        self.tables: list[list[list[tuple[str, int, int]]]] = []
        self.captions: list[str] = []
        self._rows: list[list[tuple[str, int, int]]] | None = None
        self._row: list[tuple[str, int, int]] | None = None
        self._parts: list[str] | None = None
        self._span = (1, 1)
        self._caption: list[str] | None = None

    @staticmethod
    def _span_of(attrs: dict[str, str | None], name: str) -> int:
        try:
            return max(1, int(str(attrs.get(name) or "1")))
        except ValueError:
            return 1

    def handle_starttag(self, tag, attrs):
        values = dict(attrs)
        if tag == "table":
            self._rows = []
            self.captions.append("")
        elif tag == "caption" and self._rows is not None:
            self._caption = []
        elif tag == "tr" and self._rows is not None:
            self._row = []
        elif tag in ("td", "th") and self._row is not None:
            self._parts = []
            self._span = (
                self._span_of(values, "colspan"),
                self._span_of(values, "rowspan"),
            )
        elif tag == "br" and self._parts is not None:
            self._parts.append(" ")

    def handle_endtag(self, tag):
        if tag == "table" and self._rows is not None:
            self.tables.append(self._rows)
            self._rows = None
        elif tag == "caption" and self._caption is not None:
            self.captions[-1] = " ".join("".join(self._caption).split())
            self._caption = None
        elif tag == "tr" and self._row is not None:
            if self._rows is not None:
                self._rows.append(self._row)
            self._row = None
        elif tag in ("td", "th") and self._parts is not None:
            text = " ".join("".join(self._parts).split())
            if self._row is not None:
                self._row.append((text, *self._span))
            self._parts = None

    def handle_data(self, data):
        if self._parts is not None:
            self._parts.append(data)
        elif self._caption is not None:
            self._caption.append(data)


def _lay_out_html_cells(rows: list[list[tuple[str, int, int]]]):
    """Place spanned cells the way a browser would: (grid, merge boxes).

    A cell held down by a rowspan above it occupies a column in this row too,
    so the next cell has to step past it. Without that, every row under a
    spanned header slides left by one column.
    """

    cells: dict[tuple[int, int], str] = {}
    merges: list[tuple[int, int, int, int]] = []
    occupied: set[tuple[int, int]] = set()
    for row_index, row in enumerate(rows):
        column_index = 0
        for text, across, down in row:
            while (row_index, column_index) in occupied:
                column_index += 1
            cells[(row_index, column_index)] = text
            for covered_row in range(row_index, row_index + down):
                for covered_column in range(column_index, column_index + across):
                    occupied.add((covered_row, covered_column))
            if across > 1 or down > 1:
                merges.append((
                    row_index,
                    column_index,
                    row_index + down,
                    column_index + across,
                ))
            column_index += across
    return _dense_grid(cells), merges


def _read_html_table_sheets(
    excel_path: str,
    max_rows_per_sheet: int,
) -> list[SheetText]:
    """An HTML table saved as a spreadsheet, one sheet per <table>."""

    text = _decode(Path(excel_path).read_bytes())
    reader = _HtmlTableReader()
    reader.feed(text)
    reader.close()

    sheets = []
    for index, rows in enumerate(reader.tables, start=1):
        grid, merges = _lay_out_html_cells(rows)
        caption = reader.captions[index - 1]
        title = caption or ("" if len(reader.tables) == 1 else f"Table {index}")
        sheets.append(_sheet_text(
            title, _propagate_merges(grid, merges), max_rows_per_sheet,
        ))
    return sheets or [SheetText("", "", 0, 0, 0)]


def _join_sheet_blocks(blocks: list[str], max_chars: int) -> str:
    """Join sheet blocks, stopping with a visible marker at the char budget."""
    text = "\n\n".join(blocks)
    if len(text) <= max_chars:
        return text
    return (
        text[:max_chars].rstrip()
        + f"\n... truncated: workbook exceeds {max_chars} characters ..."
    )


def _sheets_to_text(sheets: list[SheetText], max_chars: int) -> str:
    """The one-string form every non-ledger caller reads."""

    blocks = [
        f"--- Sheet: {sheet.sheet_name} (trống) ---" if not sheet.row_count
        else _sheet_header(sheet.sheet_name, sheet.row_count,
                           sheet.column_count, sheet.duplicates_dropped)
        + f"\n{sheet.content}"
        for sheet in sheets
    ]
    return _join_sheet_blocks(blocks, max_chars)


def extract_excel_text(
    excel_path: str,
    max_rows_per_sheet: int = SHEET_ROW_CAP,
    max_chars: int = 100000,
) -> str:
    """Extract XLS/XLSX workbook content as LLM-readable text grouped by sheet."""

    return _sheets_to_text(read_sheets(excel_path, max_rows_per_sheet), max_chars)


def extract_xml_text(file_path: str, max_chars: int = 120_000) -> str:
    """Flatten an XML file to readable text."""

    import xml.etree.ElementTree as ET

    raw = Path(file_path).read_text(encoding="utf-8-sig", errors="replace")
    start = raw.find("<?xml")
    if start > 0:
        raw = raw[start:]
    try:
        root = ET.fromstring(raw)
    except ET.ParseError as exc:
        return f"[XML không đọc được: {exc}]\n\n{raw[:max_chars]}"

    lines: list[str] = []

    def walk(node, path: str) -> None:
        tag = node.tag.split("}")[-1]
        here = f"{path}/{tag}" if path else tag
        text = (node.text or "").strip()
        if text:
            lines.append(f"{here}: {text}")
        for child in node:
            walk(child, here)

    walk(root, "")
    return "\n".join(lines)[:max_chars]


def extract_text_file(file_path: str, max_chars: int = 100_000) -> str:
    """A plain-text or Markdown file, decoded as best it can be."""

    return _decode(Path(file_path).read_bytes())[:max_chars]


def _pptx_shape_lines(shape) -> list[str]:
    """One shape's text: its paragraphs, or a table rendered as TSV.

    Groups are walked into, because a slide built by dragging a table and its
    title into a group hands back a single shape with no text of its own.
    """

    nested = getattr(shape, "shapes", None)
    if nested is not None:
        lines: list[str] = []
        for inner in nested:
            lines += _pptx_shape_lines(inner)
        return lines
    if getattr(shape, "has_table", False):
        return [
            "\t".join(cell.text.strip() for cell in row.cells)
            for row in shape.table.rows
        ]
    if getattr(shape, "has_text_frame", False):
        return [
            text
            for text in (
                paragraph.text.strip() for paragraph in shape.text_frame.paragraphs
            )
            if text
        ]
    return []


def extract_pptx_text(pptx_path: str, max_chars: int = 100_000) -> str:
    """Every slide of a presentation: its text, then any table it carries."""

    from pptx import Presentation

    blocks = []
    for index, slide in enumerate(Presentation(pptx_path).slides, start=1):
        lines: list[str] = []
        for shape in slide.shapes:
            lines += _pptx_shape_lines(shape)
        blocks.append(
            f"--- Slide {index} ---\n" + ("\n".join(lines) or "(no text)")
        )
    return "\n\n".join(blocks)[:max_chars]


def extract_document_text(
    file_path: str,
    ocr_timeout_seconds: float | None = None,
) -> str:
    """Extract text from supported uploaded documents."""
    
    spreadsheet_extensions = {".csv", ".xls", ".xlsx"}

    extension = Path(file_path).suffix.lower()
    if extension == ".pdf":
        return ocr_pdf(file_path, timeout_seconds=ocr_timeout_seconds)
    if extension == ".csv":
        return extract_csv_text(file_path)
    if extension in spreadsheet_extensions - {".csv"}:
        return extract_excel_text(file_path)
    if extension == ".xml":
        return extract_xml_text(file_path)
    if extension in {".txt", ".md"}:
        return extract_text_file(file_path)
    if extension == ".pptx":
        return extract_pptx_text(file_path)
    raise ValueError(f"Unsupported file extension: {extension}")
